import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from transcript import read_transcript, export_transcript, validate_course
from openpyxl import Workbook, load_workbook
from cgpa_calculator import (calculate_cgpa, calculate_semester_gpa,
                             total_completed_credits, local_ai_reply,
                             generate_ai_reply, suggest_cgpa_plans)
from streamlit.testing.v1 import AppTest

ROOT = Path(__file__).resolve().parent


def workbook_bytes(rows, extra=None):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Results"
    for row in rows:
        sheet.append(row)
    if extra:
        other = workbook.create_sheet("Other")
        for row in extra:
            other.append(row)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


class ImportTests(unittest.TestCase):
    def test_every_supplied_row_matches_source(self):
        source = ROOT / "Student_Transcript_Data (1).xlsx"
        record = read_transcript(source)
        courses = record["semesters"][0]["Courses"]
        workbook = load_workbook(source, data_only=True)
        rows = list(workbook.active.values)[1:]
        workbook.close()
        self.assertEqual(len(courses), 53)
        self.assertEqual(len(courses), len(rows))
        for number, (course, row) in enumerate(zip(courses, rows), 2):
            self.assertEqual((course["code"], course["original_credits"], course["grade"], course["grade_points"]), row)
            self.assertEqual(course["source_row"], number)
        self.assertAlmostEqual(calculate_cgpa(record["semesters"]), 371.3 / 141)
        self.assertEqual(total_completed_credits(record["semesters"]), 137)
        self.assertEqual(sum(c["code"] == "LAN 11" for c in courses), 2)
        self.assertEqual(record["reports"][0]["Courses imported"], 53)

    def test_all_sheets_and_optional_term(self):
        record = read_transcript(workbook_bytes(
            [["Title"], ["Course Code", "Credits", "Grade", "Semester"],
             [" X ", 3, "a", "Fall"], ["Y", 2, "B", None]],
            [["Code", "Units", "Letter Grade"], ["Z", 1, "C"]]))
        self.assertEqual(len(record["sheets"]), 2)
        self.assertEqual([s["Semester"] for s in record["semesters"]], ["Fall", "Other"])
        self.assertEqual(sum(len(s["Courses"]) for s in record["semesters"]), 3)

    def test_extra_sheet_retained_and_reported(self):
        record = read_transcript(workbook_bytes(
            [["Code", "Credits", "Grade"], ["X", 3, "A"]], [["Notes"], ["Keep me"]]))
        self.assertEqual(record["sheets"]["Other"][1][0], "Keep me")
        self.assertEqual(record["reports"][1]["Courses imported"], 0)

    def test_invalid_rows_are_not_silently_dropped(self):
        for credits, grade in [(-1, "A"), ("unknown", "B"), (3, "Z"), (None, "A")]:
            with self.subTest(credits=credits, grade=grade), self.assertRaisesRegex(ValueError, "row 3"):
                read_transcript(workbook_bytes([["Code", "Credits", "Grade"], ["Valid", 3, "A"], ["Bad", credits, grade]]))

    def test_pass_missing_credits_preserved(self):
        record = read_transcript(workbook_bytes([["Code", "Credits", "Grade"], ["P1", "N/A", "P"]]))
        self.assertEqual(record["semesters"][0]["Courses"][0]["credits"], 0)
        self.assertIsNone(calculate_cgpa(record["semesters"]))
        self.assertTrue(any("pass course" in w for w in record["warnings"]))

    def test_no_table_and_corrupt_file_rejected(self):
        for data in [b"invalid", workbook_bytes([["Notes"]])]:
            with self.assertRaises(Exception):
                read_transcript(data)

    def test_export_reimport_preserves_all_attempts(self):
        record = read_transcript(ROOT / "Student_Transcript_Data (1).xlsx")
        restored = read_transcript(export_transcript(record["semesters"]))
        self.assertEqual(len(restored["semesters"][0]["Courses"]), 53)
        self.assertEqual(calculate_cgpa(restored["semesters"]), calculate_cgpa(record["semesters"]))

    def test_formula_like_user_text_exported_as_text(self):
        semesters = [{"Semester": "=1+1", "Courses": [validate_course("=2+2", 3, "A")]}]
        restored = read_transcript(export_transcript(semesters))
        self.assertEqual(restored["semesters"][0]["Courses"][0]["code"], "=2+2")

    def test_validation_nan_and_unknown_grade(self):
        for credits in [float("nan"), float("inf"), -1]:
            with self.assertRaises(ValueError):
                validate_course("X", credits, "A")
        with self.assertRaises(ValueError):
            validate_course("X", 3, "Q")


class CalculationTests(unittest.TestCase):
    def test_weighted_gpa_pass_and_failure(self):
        courses = [validate_course("A", 3, "A"), validate_course("B", 1, "F"), validate_course("P", 8, "P")]
        self.assertEqual(calculate_semester_gpa(courses), 3)
        self.assertEqual(total_completed_credits([{"Courses": courses}]), 11)

    def test_empty_assistant_and_network_fallback(self):
        self.assertIn("no GPA", local_ai_reply("gpa", []))
        with patch("cgpa_calculator.generate_content", side_effect=RuntimeError("Offline")):
            with self.assertRaises(RuntimeError):
                generate_ai_reply("gpa", [], "test-key")

    def test_plan_lists_every_course_used_in_projection(self):
        courses = [validate_course(f"X{i}", 3, "D") for i in range(8)]
        plans = suggest_cgpa_plans([{"Courses": courses}], 4)
        for course in courses:
            self.assertIn(course["code"], plans[0]["note"])

    def test_offline_goal_question_routes_to_planning(self):
        semesters = [{"Courses": [validate_course("X", 3, "C")]}]
        answer = local_ai_reply("How can I reach a 3.5 CGPA?", semesters)
        self.assertIn("3.50", answer)
        self.assertIn("replaced", answer)


class InterfaceTests(unittest.TestCase):
    def setUp(self):
        self.app = AppTest.from_file(str(ROOT / "cgpa_calculator.py")).run(timeout=30)

    def click(self, label):
        next(button for button in self.app.button if button.label == label).click().run()
        self.assertFalse(self.app.exception)

    def test_initial_load_save_and_restore(self):
        self.assertFalse(self.app.exception)
        self.assertEqual(self.app.metric[1].value, "53")
        self.click("Save course changes")
        self.assertEqual(self.app.metric[1].value, "53")
        self.app.session_state.semesters[0]["Courses"].pop()
        self.app.run()
        self.assertEqual(self.app.metric[1].value, "52")
        self.click("Restore imported record")
        self.assertEqual(self.app.metric[1].value, "53")

    def test_editor_changes_refresh_metrics_and_preserve_rows(self):
        self.app.session_state["editor_1_0"] = {
            "edited_rows": {0: {"Grade": "A"}}, "added_rows": [], "deleted_rows": []}
        self.click("Save course changes")
        self.assertEqual(self.app.metric[0].value, "2.648")
        self.assertEqual(self.app.session_state.semesters[0]["Courses"][0]["source_row"], 2)
        self.app.session_state["editor_2_0"] = {
            "edited_rows": {}, "added_rows": [{"Course code": "NEW 101", "Credits": 3, "Grade": "B"}],
            "deleted_rows": [0]}
        self.click("Save course changes")
        courses = self.app.session_state.semesters[0]["Courses"]
        self.assertEqual(len(courses), 53)
        self.assertEqual(courses[0]["code"], "CSE 464")
        self.assertEqual(courses[-1]["code"], "NEW 101")

    def test_add_semester_validation_and_navigation(self):
        self.click("Create semester")
        self.assertTrue(self.app.error)
        next(t for t in self.app.text_input if t.label == "Semester name").set_value("Fall 2026")
        self.click("Create semester")
        self.assertEqual(len(self.app.session_state.semesters), 2)
        self.app.selectbox[0].select(1).run()
        self.assertFalse(self.app.exception)
        self.app.radio[0].set_value("AI Assistant").run()
        self.assertFalse(self.app.exception)
        self.app.session_state.gemini_api_key = ""
        self.app.chat_input[0].set_value("What is my GPA?").run()
        self.assertFalse(self.app.exception)
        self.assertIn("2.63", self.app.session_state.chat_history[-1]["assistant"])

    def test_quick_add_undo_and_selected_semester(self):
        next(t for t in self.app.text_input if t.label == "Semester name").set_value("Next term")
        self.click("Create semester")
        self.assertEqual(self.app.session_state.selected_group, 1)
        next(t for t in self.app.text_input if t.label == "New course code").set_value("NEW 201")
        self.click("Add course")
        self.assertEqual(self.app.session_state.selected_group, 1)
        self.assertEqual(self.app.metric[1].value, "54")
        self.click("Undo last course change")
        self.assertEqual(self.app.metric[1].value, "53")
        self.assertEqual(len(self.app.session_state.semesters), 2)

    def test_search_does_not_filter_or_delete_editor_rows(self):
        next(t for t in self.app.text_input if t.label == "Search course codes").set_value("LAN 11").run()
        self.click("Save course changes")
        self.assertEqual(self.app.metric[1].value, "53")

    def test_suggestion_fills_question_without_sending(self):
        self.app.radio[0].set_value("AI Assistant").run()
        self.click("Plan for 3.0")
        self.assertEqual(self.app.chat_input[0].value, "How can I reach a 3.0 CGPA?")
        self.assertEqual(len(self.app.session_state.chat_history), 0)

    def test_chat_followups_order_and_clear(self):
        self.app.radio[0].set_value("AI Assistant").run()
        self.app.chat_input[0].set_value("Explain my GPA").run()
        self.app.chat_input[0].set_value("Which courses need attention?").run()
        self.assertFalse(self.app.exception)
        history = self.app.session_state.chat_history
        self.assertEqual([h["user"] for h in history], ["Explain my GPA", "Which courses need attention?"])
        self.assertTrue(all(h["sent_at"] and h["answered_at"] for h in history))
        self.click("Clear chat")
        self.assertEqual(self.app.session_state.chat_history, [])

    def test_failed_message_retry_does_not_duplicate_user_message(self):
        from gemini_client import GeminiError
        self.app.radio[0].set_value("AI Assistant").run()
        next(t for t in self.app.text_input if t.label == "Gemini API key").set_value("test-key").run()
        with patch("gemini_client.urllib.request.urlopen", side_effect=GeminiError("Offline")):
            self.app.chat_input[0].set_value("Help me plan").run()
        self.assertFalse(self.app.exception)
        self.assertIn("error", self.app.session_state.chat_history[0])
        next(t for t in self.app.text_input if t.label == "Gemini API key").set_value("").run()
        self.click("Retry message")
        self.assertEqual(len(self.app.session_state.chat_history), 1)
        self.assertNotIn("error", self.app.session_state.chat_history[0])
        self.assertTrue(self.app.session_state.chat_history[0]["assistant"])

    def test_old_connection_errors_are_invalidated_after_client_update(self):
        self.app.session_state.gemini_client_revision = 1
        self.app.session_state.gemini_connection_result = (False, "Old error")
        self.app.session_state.chat_history = [{"user": "Hi", "assistant": "", "error": "Old error"}]
        self.app.run()
        self.assertFalse(self.app.exception)
        self.assertNotIn("gemini_connection_result", self.app.session_state)
        self.assertIn("Retry message", self.app.session_state.chat_history[0]["error"])

    def test_gemini_connection_result_persists_and_key_change_clears_it(self):
        with patch("gemini_client.urllib.request.urlopen", side_effect=RuntimeError("offline")):
            next(t for t in self.app.text_input if t.label == "Gemini API key").set_value("test-key").run()
            self.click("Test connection")
        self.app.run()
        self.assertIsNotNone(self.app.session_state.gemini_connection_result)
        next(t for t in self.app.text_input if t.label == "Gemini API key").set_value("different-key").run()
        self.assertNotIn("gemini_connection_result", self.app.session_state)


if __name__ == "__main__":
    unittest.main()
